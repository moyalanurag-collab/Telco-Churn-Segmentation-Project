"""Step 5 - assemble the written report (Word) from the computed results."""
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from common import PROC, FIG, ROOT

df = pd.read_csv(PROC / "telco_scored.csv")
seg = pd.read_csv(PROC / "segment_summary.csv")
comp = pd.read_csv(PROC / "model_comparison.csv")
ms = json.load(open(PROC / "model_summary.json"))
assoc = pd.read_csv(PROC / "categorical_churn_association.csv")
play = pd.read_csv(ROOT / "outputs" / "powerbi" / "dim_segments.csv")

n, churn = len(df), df.ChurnFlag.mean()
lost = df.loc[df.ChurnFlag == 1, "MonthlyCharges"].sum()
rev = df.MonthlyCharges.sum()
hr = seg[seg.churn_rate > 1.4 * churn]
flag = df[(df.predicted_churn == 1) & (df.ChurnFlag == 0)]
lift = ms["net_benefit_model"] / ms["net_benefit_all"] - 1

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)


def table(frame, widths=None):
    t = doc.add_table(rows=1, cols=len(frame.columns))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(frame.columns):
        t.rows[0].cells[i].text = str(c)
    for _, r in frame.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    doc.add_paragraph()


def fig(name, caption, w=6.3):
    doc.add_picture(str(FIG / name), width=Inches(w))
    p = doc.add_paragraph(caption)
    p.runs[0].italic = True


def bullets(items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


doc.add_heading("Customer Churn & Segmentation Analysis", 0)
doc.add_paragraph("Telco customer base - CRM retention analytics | Python, Orange Data Mining, Power BI")
doc.add_paragraph("Prepared by: Anurag Moyal")

doc.add_heading("1. Executive summary", 1)
bullets([
    f"{churn:.1%} of the {n:,} customers churned, taking ${lost:,.0f} of ${rev:,.0f} monthly revenue ({lost / rev:.1%}) with them.",
    f"Segmentation (RFM-style, K-Means) found {len(hr)} high-risk segments - {', '.join(hr.Segment)} - holding "
    f"{hr.customers.sum() / n:.0%} of customers but {hr.share_of_churn.sum():.0%} of churn and {hr.share_of_revenue_lost.sum():.0%} of lost revenue.",
    f"A gradient-boosting model separates churners from stayers with AUC {ms['oof_auc_all']:.2f} (5-fold out-of-fold). "
    f"The riskiest 20% of customers contain {ms['gains_top20']:.0%} of churners (lift {ms['gains_top20'] / .2:.1f}x); the riskiest 30% contain {ms['gains_top30']:.0%}.",
    f"Targeting customers above the {ms['threshold']:.2f} risk threshold yields about {lift:.0%} more net benefit than contacting everyone "
    f"(test set, under stated assumptions); it flags {len(flag):,} still-active customers worth ${flag.annual_revenue.sum():,.0f} annual revenue.",
    "Biggest levers: move month-to-month customers to 1-2 year contracts, fix the first-year fibre experience, bundle security/tech support, "
    "and migrate electronic-check payers to auto-pay.",
])

doc.add_heading("2. Data and cleaning", 1)
doc.add_paragraph(
    f"The IBM Telco Customer Churn dataset: {n:,} customers, 21 attributes (demographics, services, contract, billing, tenure, churn flag). "
    "No duplicates. The 11 blank TotalCharges values all belong to customers with tenure 0 (just joined) and were set to 0. "
    "'No internet/phone service' values were merged into 'No' for add-on flags, and NumServices (count of products held) was derived.")
fig("01_churn_split.png", "Figure 1 - Churn split", 3.6)

doc.add_heading("3. What drives churn", 1)
doc.add_paragraph("Association strength with churn (Cramer's V, all chi-square tests significant except gender and phone service):")
a = assoc.head(8)[["feature", "cramers_v", "highest_churn_level", "highest_rate", "lowest_churn_level", "lowest_rate"]].copy()
a["cramers_v"] = a.cramers_v.round(2)
a["highest_rate"] = (a.highest_rate * 100).round(0).astype(int).astype(str) + "%"
a["lowest_rate"] = (a.lowest_rate * 100).round(0).astype(int).astype(str) + "%"
table(a)
fig("02_churn_by_category.png", "Figure 2 - Churn rate by key attributes")
fig("04_churn_by_tenure.png", "Figure 3 - Churn falls steeply with tenure", 4.8)
bullets([
    "Contract is the strongest driver: month-to-month churns at 43% versus 3% on two-year contracts.",
    "Fibre-optic customers churn at 42% versus 7% for customers with no internet - a price/quality issue worth investigating.",
    "Electronic-check payers churn at 45% versus 15% for automatic card payers.",
    "Customers without Tech Support or Online Security churn about twice as often as those with them.",
    "Gender and phone service have no meaningful effect.",
    "Churners have a median tenure of 10 months versus 38 for stayers - risk is concentrated in the first year.",
])
fig("05_contract_internet_heatmap.png", "Figure 4 - Contract x internet service", 4.8)

doc.add_heading("4. Customer segmentation", 1)
doc.add_paragraph(
    "Because the data has no purchase history, RFM was adapted: Recency = tenure (relationship length), Frequency = number of services held, "
    "Monetary = monthly charge. Features were standardised and clustered with K-Means. k=5 was chosen: silhouette scores are close for k=3-6 "
    "(about 0.39-0.43), and k=5 is the smallest solution that separates two distinct high-risk groups that k=4 merges.")
fig("07_kmeans_k_selection.png", "Figure 5 - Choosing k", 5.5)
s = seg[["Segment", "customers", "avg_tenure_months", "avg_services", "avg_monthly_charge", "churn_rate", "share_of_churn", "share_of_revenue_lost"]].copy()
for c in ["avg_tenure_months", "avg_services", "avg_monthly_charge"]:
    s[c] = s[c].round(1)
for c in ["churn_rate", "share_of_churn", "share_of_revenue_lost"]:
    s[c] = (s[c] * 100).round(0).astype(int).astype(str) + "%"
s.columns = ["Segment", "Customers", "Tenure (m)", "Services", "Monthly $", "Churn", "% of churn", "% of rev lost"]
table(s)
fig("10_segment_churn_overview.png", "Figure 6 - Segment churn, size and revenue lost")
fig("09_segment_rfm_heatmap.png", "Figure 7 - RFM-style profile (std. deviations from average)", 4.8)

doc.add_heading("5. Churn prediction model", 1)
doc.add_paragraph("Four models were trained on a stratified 80% split with 5-fold CV and evaluated on the held-out 20%:")
c = comp[["model", "cv_auc_mean", "test_auc", "test_pr_auc", "brier"]].round(3)
c.columns = ["Model", "CV AUC", "Test AUC", "Test PR-AUC", "Brier"]
table(c)
doc.add_paragraph(
    "All four models perform within about 0.005 AUC of each other, a sign that the signal is mostly captured by contract, tenure and internet service. "
    f"Gradient boosting had the best CV AUC and best probability calibration (lowest Brier score) and was selected; "
    "its out-of-fold probabilities score every customer in the dashboard, so reported risk is not inflated by training on the same rows.")
fig("11_roc_pr_curves.png", "Figure 8 - ROC and precision-recall curves (test set)")
fig("17_permutation_importance.png", "Figure 9 - What the model relies on", 5)
fig("16_logistic_drivers.png", "Figure 10 - Logistic regression drivers (interpretable cross-check)", 5.5)
fig("14_gains_curve.png", "Figure 11 - Cumulative gains", 4.5)

doc.add_heading("6. Turning predictions into decisions", 1)
a_ = ms["assumptions"]
doc.add_paragraph(
    f"A default 0.5 cut-off is wrong for retention: missing a churner costs far more than a wasted call. Assuming ${a_['contact_cost']:.0f} per outreach, "
    f"{a_['offer_success']:.0%} acceptance among would-be churners and {a_['months_saved']} months of retained revenue, the net-benefit-maximising threshold "
    f"(chosen on training out-of-fold scores, not the test set) is {ms['threshold']:.2f}. On the test set this gives precision {ms['test_precision']:.0%} and recall "
    f"{ms['test_recall']:.0%}, and net benefit ${ms['net_benefit_model']:,.0f} versus ${ms['net_benefit_all']:,.0f} for contacting everyone. "
    "These assumptions are placeholders; replace them with the company's real figures.")
fig("12_profit_curve.png", "Figure 12 - Net benefit by threshold", 4.8)
fig("13_confusion_matrix.png", "Figure 13 - Confusion matrix at the chosen threshold", 3.4)

doc.add_heading("7. Who to target and what to do", 1)
p = play[["priority", "Segment", "recommended_programme", "recommended_action"]]
p.columns = ["Priority", "Segment", "Programme", "Action"]
table(p)

doc.add_heading("8. Limitations and next steps", 1)
bullets([
    "Snapshot data: no time dimension, so real recency and time-to-churn (survival analysis) cannot be modelled.",
    "Business assumptions (cost, acceptance rate, months saved) are illustrative; validate with finance/CRM.",
    "Findings are associations. Run an A/B test of the retention offers (control group vs treated) to measure true uplift, ideally with an uplift model.",
    "Next: add usage, complaint and support-ticket data; monitor model drift monthly; refresh the target list weekly.",
])
doc.add_heading("Resume line", 2)
doc.add_paragraph(
    f"Analysed {n:,} telecom customers (Python, Orange, Power BI): identified {len(hr)} high-risk segments driving {hr.share_of_churn.sum():.0%} of churn "
    f"and {hr.share_of_revenue_lost.sum():.0%} of lost revenue; built a churn model (AUC {ms['oof_auc_all']:.2f}) capturing {ms['gains_top20']:.0%} of churners in the top 20% of customers "
    "and proposed segment-specific retention actions.")

out = ROOT / "report"
out.mkdir(exist_ok=True)
doc.save(out / "Churn_Analysis_Report.docx")
print("Report saved")
